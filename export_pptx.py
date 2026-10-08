#!/usr/bin/python3.14
# PPTX + ODP Export: Screenshots jeder Folie + Video eingebettet mit Autoplay
# Ausfuehren: python3.14 export_pptx.py
import os
import subprocess
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

hier = os.path.dirname(os.path.abspath(__file__))
os.makedirs('/tmp/folien', exist_ok=True)

# === 1. Screenshots aller Folien via Playwright ===
print('Screenshots erstellen...')
with sync_playwright() as p:
    b = p.chromium.launch(executable_path='/usr/bin/google-chrome-stable')
    pg = b.new_page(viewport={'width': 1920, 'height': 1080})
    pg.goto('file://' + os.path.join(hier, 'index.html'))
    anzahl = pg.locator('.slide').count()
    bilder, video_folie = [], None
    for i in range(anzahl):
        pg.wait_for_timeout(900)
        if pg.locator('.slide.active video').count():
            video_folie = i
            pg.evaluate('v => { v.pause(); v.currentTime = v.duration / 2; }',
                        pg.locator('.slide.active video').element_handle())
            pg.wait_for_timeout(500)
        pfad = f'/tmp/folien/{i:02d}.png'
        pg.screenshot(path=pfad)
        bilder.append(pfad)
        pg.keyboard.press('ArrowRight')
    b.close()
print(f'{len(bilder)} Folien fotografiert, Video auf Folie {video_folie + 1 if video_folie is not None else "-"}')

# === 2. Video vorbereiten (stumm, Poster) ===
VIDEO = os.path.join(hier, 'spaghetti.mp4')
STUMM = '/tmp/folien/spaghetti-stumm.mp4'
POSTER = '/tmp/folien/poster.png'

if os.path.exists(VIDEO):
    print('Video vorbereiten...')
    subprocess.run(['ffmpeg', '-loglevel', 'error', '-y', '-i', VIDEO, '-an', '-c:v', 'copy', STUMM], check=True)
    subprocess.run(['ffmpeg', '-loglevel', 'error', '-y', '-ss', '1', '-i', VIDEO, '-frames:v', '1', POSTER], check=True)
    info = subprocess.run(['ffprobe', '-v', 'error', '-select_streams', 'v:0', '-show_entries',
                           'stream=width,height,duration', '-of', 'csv=p=0', VIDEO],
                          capture_output=True, text=True, check=True).stdout.strip().split(',')
    VIDEO_W, VIDEO_H, DAUER_MS = int(info[0]), int(info[1]), int(float(info[2]) * 1000)
    HAS_VIDEO = True
else:
    print('WARNUNG: spaghetti.mp4 nicht gefunden, Video-Folie wird Screenshot')
    HAS_VIDEO = False


def autoplay(folie, spid):
    xml = f'''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{DAUER_MS}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    alt = folie._element.find(qn('p:timing'))
    if alt is not None:
        folie._element.remove(alt)
    folie._element.append(etree.fromstring(xml))


# === 3. PPTX bauen ===
print('PPTX erstellen...')
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9

for i, bild in enumerate(bilder):
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    if i == video_folie and HAS_VIDEO:
        folie.background.fill.solid()
        folie.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
        breite = prs.slide_width
        hoehe = int(breite * VIDEO_H / VIDEO_W)
        video = folie.shapes.add_movie(STUMM, 0, (prs.slide_height - hoehe) // 2, breite, hoehe,
                                       poster_frame_image=POSTER, mime_type='video/mp4')
        autoplay(folie, video.shape_id)
    else:
        folie.shapes.add_picture(bild, 0, 0, prs.slide_width, prs.slide_height)

pptx_path = os.path.join(hier, 'ethik-praesentation.pptx')
prs.save(pptx_path)
print(f'PPTX gespeichert: {pptx_path}')

# === 4. ODP via LibreOffice ===
print('ODP konvertieren...')
subprocess.run(['libreoffice', '--headless', '--convert-to', 'odp', pptx_path, '--outdir', hier],
               capture_output=True, check=True)
odp_path = os.path.join(hier, 'ethik-praesentation.odp')
print(f'ODP gespeichert: {odp_path}')

print(f'\nFertig: {len(bilder)} Folien, Video auf Folie {video_folie + 1 if video_folie is not None else "-"}')
