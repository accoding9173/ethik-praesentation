# Backup-PPTX erzeugen: jede Folie der Website als Bild, das Video als echtes Video eingebettet
# Ausführen: python make_pptx.py (braucht playwright + python-pptx)
import os
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

hier = os.path.dirname(os.path.abspath(__file__))
os.makedirs("/tmp/folien", exist_ok=True)

with sync_playwright() as p:
    b = p.chromium.launch(executable_path="/usr/bin/google-chrome-stable")
    pg = b.new_page(viewport={"width": 1920, "height": 1080})
    pg.goto("file://" + os.path.join(hier, "index.html"))
    anzahl = pg.locator(".slide").count()
    bilder, video_folie = [], None
    for i in range(anzahl):
        pg.wait_for_timeout(900)  # Überblendung abwarten
        if pg.locator(".slide.active video").count():
            video_folie = i
            # Standbild aus der Mitte als Vorschaubild
            pg.evaluate("v => { v.pause(); v.currentTime = v.duration / 2; }",
                        pg.locator(".slide.active video").element_handle())
            pg.wait_for_timeout(500)
        pfad = f"/tmp/folien/{i:02d}.png"
        pg.screenshot(path=pfad)
        bilder.append(pfad)
        pg.keyboard.press("ArrowRight")
    b.close()

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
for i, bild in enumerate(bilder):
    folie = prs.slides.add_slide(prs.slide_layouts[6])
    if i == video_folie:
        # schwarzer Hintergrund, Video zentriert in 16:9 über die ganze Höhe
        folie.background.fill.solid()
        from pptx.dml.color import RGBColor
        folie.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
        folie.shapes.add_movie(os.path.join(hier, "spaghetti.mp4"), 0, 0,
                               prs.slide_width, prs.slide_height,
                               poster_frame_image=bild, mime_type="video/mp4")
    else:
        folie.shapes.add_picture(bild, 0, 0, prs.slide_width, prs.slide_height)

prs.save(os.path.join(hier, "praesentation-backup.pptx"))
print(len(bilder), "Folien, Video auf Folie", video_folie + 1 if video_folie is not None else "-")
